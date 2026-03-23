"""
db-writer: export_excel.py
Export the full LNP database to a multi-sheet Excel workbook.
Called on demand via POST /export or scheduled job.
"""

import json
import logging
import os
import sqlite3
import sys
from datetime import datetime

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
logger = logging.getLogger("export-excel")

TABLES = [
    "Material",
    "Formulation",
    "FormulationComponent",
    "Batch",
    "AssayResults_Physical",
    "InVivo_study",
    "InVivo_FLUC",
    "InVivo_EPO",
    "Toxicity",
]


def _open_db(db_path: str) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path, timeout=10)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    conn.row_factory = sqlite3.Row
    return conn


def export_excel(db_path: str, exports_dir: str) -> str:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.error("EXPORT_DEPENDENCY_MISSING: openpyxl not installed")
        raise RuntimeError("EXPORT_DEPENDENCY_MISSING: install openpyxl to use export feature")

    if not os.path.exists(db_path):
        raise FileNotFoundError(f"Database not found: {db_path}")

    conn = _open_db(db_path)
    wb = openpyxl.Workbook()
    first_sheet = True

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
    header_align = Alignment(horizontal="center")

    try:
        for table in TABLES:
            try:
                rows = conn.execute(f"SELECT * FROM {table}").fetchall()  # noqa: S608
            except sqlite3.OperationalError:
                logger.warning("Table %s not found in DB; skipping", table)
                continue

            if first_sheet:
                ws = wb.active
                ws.title = table
                first_sheet = False
            else:
                ws = wb.create_sheet(title=table)

            if not rows:
                ws.append([table, "(no data)"])
                continue

            # Header row
            col_names = list(rows[0].keys())
            ws.append(col_names)
            for col_idx, _ in enumerate(col_names, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            # Data rows
            for row in rows:
                ws.append(list(row))

            # Auto-width
            for col_cells in ws.columns:
                max_len = max(
                    (len(str(c.value)) if c.value is not None else 0 for c in col_cells),
                    default=10,
                )
                ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)

    finally:
        conn.close()

    os.makedirs(exports_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(exports_dir, f"export_{ts}.xlsx")
    wb.save(out_path)
    logger.info("Exported database to %s", out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: export_excel.py <db_path> <exports_dir>")
        sys.exit(1)
    print(export_excel(sys.argv[1], sys.argv[2]))
